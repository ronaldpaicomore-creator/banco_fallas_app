"""
drive_indexer.py
-----------------
Recorre recursivamente una carpeta raíz de Google Drive (incluye Unidades
compartidas), descarga/exporta cada archivo soportado, extrae su texto y
devuelve una lista de objetos langchain Document listos para indexar.

Estructura esperada en el Drive (ejemplo):
    Banco de Fallas/
        Linea 1/
            Electrica/
                archivo1.pptx
                archivo2.pdf
            Mecanica/
                archivo3.docx
        Linea 2/
            ...

No importa cuántos niveles de subcarpetas tengas: la función es recursiva,
así que da igual si hay 2 o 5 niveles.
"""

import io
import os
from typing import List

import streamlit as st
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from langchain_core.documents import Document

from pypdf import PdfReader
import docx
from pptx import Presentation

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
CREDENTIALS_FILE = "credentials.json"

# Tipos de archivo que sabemos leer
MIME_HANDLERS = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    # Archivos nativos de Google (se exportan, no se descargan directo)
    "application/vnd.google-apps.document": "gdoc",
    "application/vnd.google-apps.presentation": "gslides",
}

FOLDER_MIME = "application/vnd.google-apps.folder"


def get_drive_service():
    """Autentica con la API de Google Drive desde Secrets (nube) o credentials.json (local)."""
    if "gcp_service_account" in st.secrets:
        creds_dict = dict(st.secrets["gcp_service_account"])
        
        # Limpia y arregla la clave privada sea cual sea el formato pegado
        if "private_key" in creds_dict:
            pk = creds_dict["private_key"]
            # Si se pegó con \n literales, los convierte a saltos de línea reales
            if "\\n" in pk:
                pk = pk.replace("\\n", "\n")
            # Elimina comillas extra o espacios accidentales al inicio/final
            creds_dict["private_key"] = pk.strip().strip('"').strip("'")
            
        creds = service_account.Credentials.from_service_account_info(
            creds_dict,
            scopes=SCOPES
        )
    else:
        creds = service_account.Credentials.from_service_account_file(
            CREDENTIALS_FILE,
            scopes=SCOPES
        )
    return build("drive", "v3", credentials=creds)


def list_children(service, folder_id: str):
    """Lista todo lo que hay dentro de una carpeta (soporta Unidades compartidas)."""
    items = []
    page_token = None
    while True:
        response = (
            service.files()
            .list(
                q=f"'{folder_id}' in parents and trashed = false",
                spaces="drive",
                corpora="allDrives",
                includeItemsFromAllDrives=True,
                supportsAllDrives=True,
                fields="nextPageToken, files(id, name, mimeType, webViewLink, modifiedTime)",
                pageToken=page_token,
                pageSize=200,
            )
            .execute()
        )
        items.extend(response.get("files", []))
        page_token = response.get("nextPageToken")
        if not page_token:
            break
    return items


def walk_drive(service, folder_id: str, path_parts=None):
    """Recorre recursivamente y devuelve una lista de dicts:
    {id, name, mimeType, webViewLink, ruta}"""
    if path_parts is None:
        path_parts = []

    found = []
    for item in list_children(service, folder_id):
        if item["mimeType"] == FOLDER_MIME:
            found.extend(walk_drive(service, item["id"], path_parts + [item["name"]]))
        else:
            item["ruta"] = " / ".join(path_parts)
            found.append(item)
    return found


def _download_bytes(service, file_id: str) -> bytes:
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buf.getvalue()


def _export_bytes(service, file_id: str, export_mime: str) -> bytes:
    request = service.files().export_media(fileId=file_id, mimeType=export_mime)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buf.getvalue()


def extract_text(service, file_meta: dict) -> str:
    """Extrae texto de un archivo según su tipo. Devuelve '' si no se puede leer."""
    mime = file_meta["mimeType"]
    file_id = file_meta["id"]
    kind = MIME_HANDLERS.get(mime)

    try:
        if kind == "pdf":
            raw = _download_bytes(service, file_id)
            reader = PdfReader(io.BytesIO(raw))
            return "\n".join((p.extract_text() or "") for p in reader.pages)

        elif kind == "docx":
            raw = _download_bytes(service, file_id)
            document = docx.Document(io.BytesIO(raw))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)

        elif kind == "pptx":
            raw = _download_bytes(service, file_id)
            prs = Presentation(io.BytesIO(raw))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)

        elif kind == "gdoc":
            raw = _export_bytes(service, file_id, "text/plain")
            return raw.decode("utf-8", errors="ignore")

        elif kind == "gslides":
            raw = _export_bytes(service, file_id, "text/plain")
            return raw.decode("utf-8", errors="ignore")

        else:
            return ""  # tipo no soportado (imágenes, hojas de cálculo, etc.)

    except Exception as e:
        print(f"[WARN] No se pudo leer '{file_meta['name']}': {e}")
        return ""


def process_drive_folder(root_folder_id: str) -> List[Document]:
    """Punto de entrada: recorre el Drive completo y devuelve Documents de LangChain."""
    service = get_drive_service()
    files = walk_drive(service, root_folder_id)
    return _procesar_lista_de_archivos(service, files)


def process_drive_folder_incremental(root_folder_id: str, ids_ya_indexados: set, pares_ya_indexados: set) -> List[Document]:
    """Como process_drive_folder, pero se salta cualquier archivo que ya
    esté indexado — ya sea porque su id de Drive coincide, o porque su
    combinación (nombre, ruta) ya existía en un índice más antiguo que no
    guardaba el id. Así solo se procesan los archivos realmente nuevos."""
    service = get_drive_service()
    files = walk_drive(service, root_folder_id)
    files_nuevos = [
        f for f in files
        if f["id"] not in ids_ya_indexados and (f["name"], f["ruta"]) not in pares_ya_indexados
    ]
    print(f"[INFO] {len(files_nuevos)} archivos nuevos de {len(files)} totales en Drive.")
    return _procesar_lista_de_archivos(service, files_nuevos)


def _procesar_lista_de_archivos(service, files) -> List[Document]:
    documents = []
    for f in files:
        if f["mimeType"] not in MIME_HANDLERS:
            continue  # saltar imágenes, hojas de cálculo, atajos, etc.

        text = extract_text(service, f)
        if not text.strip():
            continue

        # La ruta suele ser algo como "Linea 1 / Electrica"
        ruta_partes = f["ruta"].split(" / ") if f["ruta"] else []
        linea = ruta_partes[0] if len(ruta_partes) > 0 else "Sin línea"
        especialidad = ruta_partes[1] if len(ruta_partes) > 1 else "Sin especialidad"

        # modifiedTime viene como "2025-06-12T15:30:00.000Z" -> nos quedamos
        # con el mes en formato "2025-06" para poder filtrar por mes.
        fecha_completa = f.get("modifiedTime", "")
        fecha_mes = fecha_completa[:7] if fecha_completa else ""

        documents.append(
            Document(
                page_content=text,
                metadata={
                    "id": f["id"],
                    "nombre": f["name"],
                    "ruta": f["ruta"],
                    "linea": linea,
                    "especialidad": especialidad,
                    "link": f.get("webViewLink", ""),
                    "tipo": MIME_HANDLERS[f["mimeType"]],
                    "fecha_mod": fecha_completa,
                    "mes": fecha_mes,
                },
            )
        )

    print(f"[INFO] {len(documents)} documentos indexados de {len(files)} archivos procesados.")
    return documents